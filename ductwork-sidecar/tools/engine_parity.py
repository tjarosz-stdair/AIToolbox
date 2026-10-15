"""Engine parity: the wrapped commercial engine must reproduce an audited reference exactly.

The engine is wrapped, not rewritten (plan section 6.7), so the guarantee that wrapping did not change
its answers is a regression test: run the engine on a reference plan set with the reference
settings, and compare its summary to the audited reference workbook, line by line.

    python tools/engine_parity.py --engine <folder containing hvac_takeoff> \\
        --plan plans.pdf --config project.yaml --expected Audited.xlsx [--workdir DIR]

No customer data lives in this repository: the plan, settings and reference are passed in as paths.
Exit code 0 means every line matched (quantities within a tolerance of one thousandth); 1 means
differences, which are printed; 2 means the engine could not be run or a file could not be read.

Rows whose Basis in the reference is "Callout count" are display-only: the audited workbook shows 0 for
them while the engine's CSV lists the raw callout number. They are set aside by default and reported as a
count; pass --include-callouts to compare them as well.
"""

from __future__ import annotations

import argparse
import csv
import os
import subprocess
import sys
import tempfile
from collections import defaultdict
from dataclasses import dataclass

TOLERANCE = 1e-3
SUMMARY_SHEET = "Summary"
HEADER = ("category", "item", "size / model", "unit", "total qty")
CALLOUT_BASIS = "callout count"


@dataclass(frozen=True)
class Line:
    category: str
    item: str
    size: str
    unit: str
    qty: float
    basis: str = ""

    @property
    def key(self) -> tuple[str, str, str, str]:
        return (_norm(self.category), _norm(self.item), _norm(self.size), _norm(self.unit))


@dataclass(frozen=True)
class Difference:
    key: tuple[str, str, str, str]
    expected: float | None
    actual: float | None

    def describe(self) -> str:
        label = " | ".join(part for part in self.key if part)
        if self.expected is None:
            return f"extra in actual:   {label}  actual={self.actual:g}"
        if self.actual is None:
            return f"missing in actual: {label}  expected={self.expected:g}"
        return f"quantity differs:  {label}  expected={self.expected:g} actual={self.actual:g}"


def _norm(value: str | None) -> str:
    return " ".join((value or "").split()).lower()


def _number(value) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def lines_from_rows(rows) -> list[Line]:
    """Turns spreadsheet or CSV rows into lines. Section titles and blank rows (no unit or quantity) are skipped."""
    lines: list[Line] = []
    header_seen = False
    for row in rows:
        cells = list(row)
        if not header_seen:
            if tuple(_norm(str(c)) for c in cells[:5]) == HEADER:
                header_seen = True
            continue
        if len(cells) < 5:
            continue
        category, item, size, unit, qty = cells[0], cells[1], cells[2], cells[3], _number(cells[4])
        if not unit or qty is None:
            continue
        basis = str(cells[6]) if len(cells) > 6 and cells[6] is not None else ""
        lines.append(Line(str(category or ""), str(item or ""), str(size or ""), str(unit), qty, basis))
    return lines


def load_expected_xlsx(path: str) -> list[Line]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if SUMMARY_SHEET not in workbook.sheetnames:
            raise ValueError(f"{path} has no '{SUMMARY_SHEET}' sheet.")
        return _carry_category(lines_from_rows(workbook[SUMMARY_SHEET].iter_rows(values_only=True)))
    finally:
        workbook.close()


def load_actual_csv(path: str) -> list[Line]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        return _carry_category([Line(*row_cells) for row_cells in _csv_lines(reader)])


def _csv_lines(reader):
    header = next(reader)
    columns = [_norm(h) for h in header]
    wanted = {"category": "category", "item": "item", "size": "size", "unit": "unit", "total_qty": "total_qty"}
    index = {k: columns.index(v) for k, v in wanted.items() if v in columns}
    for row in reader:
        qty = _number(row[index["total_qty"]]) if "total_qty" in index and len(row) > index["total_qty"] else None
        unit = row[index["unit"]] if "unit" in index and len(row) > index["unit"] else ""
        if not unit or qty is None:
            continue
        yield (row[index["category"]], row[index["item"]], row[index.get("size", index["item"])] if "size" in index else "", unit, qty)


def _carry_category(lines: list[Line]) -> list[Line]:
    # A workbook can leave the category blank on lines that continue the previous one.
    result, current = [], ""
    for line in lines:
        current = line.category or current
        result.append(Line(current, line.item, line.size, line.unit, line.qty, line.basis))
    return result


def split_callouts(expected: list[Line]) -> tuple[list[Line], set[tuple]]:
    """Separates reference lines that are display-only callouts. Returns the rest and the keys set aside."""
    callout_keys = {l.key for l in expected if _norm(l.basis) == CALLOUT_BASIS}
    return [l for l in expected if l.key not in callout_keys], callout_keys


def compare(expected: list[Line], actual: list[Line], tolerance: float = TOLERANCE, ignore_keys: set[tuple] | None = None) -> list[Difference]:
    """Compares per (category, item, size, unit). Lines that share a key are summed, so ordering and splitting do not matter."""
    ignore_keys = ignore_keys or set()
    expected = [l for l in expected if l.key not in ignore_keys]
    actual = [l for l in actual if l.key not in ignore_keys]

    def totals(lines: list[Line]) -> dict[tuple, float]:
        sums: dict[tuple, float] = defaultdict(float)
        for line in lines:
            sums[line.key] += line.qty
        return sums

    e, a = totals(expected), totals(actual)
    differences = []
    for key in sorted(set(e) | set(a)):
        if key not in a:
            differences.append(Difference(key, e[key], None))
        elif key not in e:
            differences.append(Difference(key, None, a[key]))
        elif abs(e[key] - a[key]) > tolerance:
            differences.append(Difference(key, e[key], a[key]))
    return differences


def run_engine(engine_dir: str, plan: str, config: str | None, out_dir: str, timeout: int = 1800) -> str:
    """Runs the engine without markup and returns the path of its summary CSV."""
    command = [sys.executable, "-m", "hvac_takeoff", plan, "-o", out_dir, "--no-markup", "--no-cache"]
    if config:
        command += ["-c", config]
    env = dict(os.environ, PYTHONPATH=engine_dir, PYTHONUTF8="1")
    done = subprocess.run(command, cwd=engine_dir, env=env, capture_output=True, text=True, timeout=timeout)
    if done.returncode != 0:
        raise RuntimeError(f"The engine failed ({done.returncode}):\n{done.stderr[-2000:]}")
    summaries = [f for f in os.listdir(out_dir) if f.endswith("_summary.csv")]
    if len(summaries) != 1:
        raise RuntimeError(f"Expected one summary CSV in {out_dir}, found {summaries}.")
    return os.path.join(out_dir, summaries[0])


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--engine", required=True, help="folder that contains the hvac_takeoff package")
    parser.add_argument("--plan", required=True)
    parser.add_argument("--config")
    parser.add_argument("--expected", required=True, help="the audited reference workbook (.xlsx)")
    parser.add_argument("--workdir", help="where the engine writes its output (default: a temporary folder)")
    parser.add_argument("--include-callouts", action="store_true", help="also compare display-only callout rows")
    args = parser.parse_args(argv)

    try:
        expected = load_expected_xlsx(args.expected)
        with tempfile.TemporaryDirectory() as tmp:
            out = args.workdir or tmp
            os.makedirs(out, exist_ok=True)
            actual = load_actual_csv(run_engine(args.engine, args.plan, args.config, out))
    except (OSError, RuntimeError, ValueError, subprocess.TimeoutExpired) as ex:
        print(f"Could not complete the parity run: {ex}", file=sys.stderr)
        return 2

    ignored: set[tuple] = set() if args.include_callouts else split_callouts(expected)[1]
    differences = compare(expected, actual, ignore_keys=ignored)
    print(f"reference lines: {len(expected)}   engine lines: {len(actual)}   "
          f"callout rows set aside: {len(ignored)}   differences: {len(differences)}")
    for d in differences[:100]:
        print("  " + d.describe())
    if len(differences) > 100:
        print(f"  ... and {len(differences) - 100} more")
    return 0 if not differences else 1


if __name__ == "__main__":
    raise SystemExit(main())
